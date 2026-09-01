"""
Build the DATA / GENERATED_AT block for fpl_tracker_dashboard.html
---------------------------------------------------------
Pulls the "FPL 3-Team Tracker" Google Sheet (must be shared as
"Anyone with the link can view") and regenerates the dashboard's
embedded DATA object, replacing the content between the
DATA_START/DATA_END and GENERATED_AT_START/GENERATED_AT_END marker
comments.

Requires: pip install openpyxl

Usage:
    python script/build_dashboard_data.py
"""

import io
import json
import os
import re
import urllib.request
from datetime import date, datetime, timezone

import openpyxl

SHEET_ID = "1w7Ih8neSVFHlLcScYvsrcwzzf4T1iCcGHmbTeYkAGEo"
EXPORT_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DASHBOARD_PATH = os.path.join(os.path.dirname(SCRIPT_DIR), "fpl_tracker_dashboard.html")


def fetch_workbook():
    req = urllib.request.Request(EXPORT_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        raw = resp.read()
    return openpyxl.load_workbook(io.BytesIO(raw), data_only=True)


def num(v):
    """Collapse whole-number floats to int (6.0 -> 6), pass through everything else."""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def text(v):
    if v in (None, ""):
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d %b %Y")
    if isinstance(v, date):
        return v.strftime("%d %b %Y")
    return v


def player_rows(ws):
    rows = []
    for row in ws.iter_rows(values_only=True):
        pos = row[1] if len(row) > 1 else None
        if pos not in ("GK", "DEF", "MID", "FWD"):
            continue
        rows.append({
            "pos": pos,
            "player": text(row[2]),
            "club": text(row[3]),
            "price": num(row[4]),
            "status": text(row[5]),
            "notes": text(row[7]),
            "captain": text(row[6]),
        })
    return rows


def build_data(wb):
    teams = [name.split("Squad - ", 1)[1] for name in wb.sheetnames if name.startswith("Squad - ")]

    weekly = {team: [] for team in teams}
    series = {team: [] for team in teams}
    ws = wb["Weekly Log"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        team, gw, points = row[0], row[1], row[2]
        if not team or points is None:
            continue
        if team not in weekly:
            continue
        weekly[team].append({
            "gw": num(gw),
            "points": num(points),
            "bench": num(row[3]),
            "transfers": num(row[4]),
            "cost": num(row[5]),
            "chip": text(row[6]),
            "captain": text(row[7]),
            "vice": text(row[8]),
            "rank": num(row[9]),
            "value": num(row[10]),
            "bank": num(row[11]),
            "net": num(row[12]),
            "notes": text(row[14]),
        })
        series[team].append({"gw": num(gw), "total": num(row[13])})

    moves = {team: [] for team in teams}
    ws = wb["Transfer Log"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        team = row[0]
        if not team or team not in moves:
            continue
        moves[team].append({
            "gw": num(row[1]),
            "date": text(row[2]),
            "out": text(row[3]),
            "in": text(row[4]),
            "priceOut": num(row[5]),
            "priceIn": num(row[6]),
            "cost": num(row[7]),
            "reason": text(row[8]),
            "outcome": text(row[9]),
        })

    squads = {team: player_rows(wb[f"Squad - {team}"]) for team in teams}

    return {
        "teams": teams,
        "weekly": weekly,
        "moves": moves,
        "squads": squads,
        "series": series,
    }


def replace_between(html, start_marker, end_marker, new_inner):
    pattern = re.compile(
        re.escape(start_marker) + r".*?" + re.escape(end_marker),
        re.DOTALL,
    )
    if not pattern.search(html):
        raise SystemExit(f"Could not find {start_marker}/{end_marker} markers in {DASHBOARD_PATH}")
    replacement = f"{start_marker}\n{new_inner}\n{end_marker}"
    return pattern.sub(lambda _: replacement, html, count=1)


def inject(data):
    with open(DASHBOARD_PATH, "r", encoding="utf-8") as f:
        html = f.read()

    data_line = f"const DATA = {json.dumps(data, ensure_ascii=False)};"
    html = replace_between(html, "/* DATA_START */", "/* DATA_END */", data_line)

    generated_at = datetime.now(timezone.utc).astimezone().strftime("%d %b %Y, %H:%M")
    gen_line = f'const GENERATED_AT = "{generated_at}";'
    html = replace_between(html, "/* GENERATED_AT_START */", "/* GENERATED_AT_END */", gen_line)

    with open(DASHBOARD_PATH, "w", encoding="utf-8") as f:
        f.write(html)

    return generated_at


def main():
    wb = fetch_workbook()
    data = build_data(wb)
    generated_at = inject(data)
    print(f"Wrote DATA for {len(data['teams'])} team(s): {', '.join(data['teams'])}")
    for team in data["teams"]:
        print(f"  {team}: {len(data['weekly'][team])} GW(s) logged, "
              f"{len(data['squads'][team])} squad player(s), "
              f"{len(data['moves'][team])} move(s)")
    print(f"GENERATED_AT set to {generated_at}")


if __name__ == "__main__":
    main()
